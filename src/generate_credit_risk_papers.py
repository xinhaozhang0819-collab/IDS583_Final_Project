from __future__ import annotations

import json
import math
import re
import shutil
from datetime import datetime
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from sklearn.metrics import brier_score_loss, roc_auc_score, roc_curve


PROJECT_ROOT = Path(__file__).resolve().parents[1]
REPO_ROOT = PROJECT_ROOT.parent
PROCESSED_DIR = PROJECT_ROOT / "data" / "processed"
REPORT_DIR = PROJECT_ROOT / "reports"
FIG_TEMPORAL = REPORT_DIR / "figures" / "temporal_model"
FIG_LOSS = REPORT_DIR / "figures" / "loss_reserve"

DOC_TARGETS = [
    PROJECT_ROOT / "output" / "doc" / "credit_risk_selected_chapters_academic.docx",
    REPO_ROOT / "credit risk paper.docx",
]


def _fmt_num(value, digits=4):
    if value is None:
        return "NA"
    try:
        if pd.isna(value):
            return "NA"
        return f"{float(value):,.{digits}f}"
    except Exception:
        return str(value)


def _fmt_pct(value, digits=1):
    if value is None:
        return "NA"
    try:
        if pd.isna(value):
            return "NA"
        return f"{float(value) * 100:.{digits}f}%"
    except Exception:
        return str(value)


def _fmt_money(value):
    if value is None or pd.isna(value):
        return "NA"
    value = float(value)
    if abs(value) >= 1_000_000_000:
        return f"${value / 1_000_000_000:.2f}B"
    if abs(value) >= 1_000_000:
        return f"${value / 1_000_000:.1f}M"
    return f"${value:,.0f}"


def _short_params(params_text):
    try:
        params = json.loads(params_text)
    except Exception:
        return str(params_text)
    keys = ["n_estimators", "max_depth", "learning_rate", "C", "class_weight", "hidden_layer_sizes"]
    parts = []
    for key in keys:
        if key in params:
            parts.append(f"{key}={params[key]}")
    return ", ".join(parts) if parts else str(params)


def _ks(y_true, y_score):
    fpr, tpr, _ = roc_curve(y_true, y_score)
    return float((tpr - fpr).max())


def _metrics(y_true, y_score):
    y_true = pd.Series(y_true).fillna(0).astype(int)
    y_score = pd.Series(y_score).fillna(0).clip(0, 1)
    return {
        "auc": float(roc_auc_score(y_true, y_score)) if y_true.nunique() > 1 else 0.5,
        "ks": _ks(y_true, y_score) if y_true.nunique() > 1 else 0.0,
        "brier": float(brier_score_loss(y_true, y_score)),
    }


def _read_inputs():
    runtime = json.loads((PROCESSED_DIR / "calendar_30min_training_runtime.json").read_text())
    stage2 = json.loads((PROCESSED_DIR / "stage2_champion_config.json").read_text())
    view_summary = json.loads(
        (PROCESSED_DIR / "calendar_survival_30min_training_view_summary.json").read_text()
    )
    search = pd.read_csv(PROCESSED_DIR / "calendar_30min_grid_search_results.csv")
    portfolio = pd.read_csv(PROCESSED_DIR / "portfolio_expected_loss_summary.csv")
    segment = pd.read_csv(PROCESSED_DIR / "segment_expected_loss_summary.csv")
    hazard_comparison = pd.read_csv(PROCESSED_DIR / "hazard_model_comparison.csv")
    predictions = pd.read_csv(PROCESSED_DIR / "test_with_pd_best_model.csv")
    charged_proxy = pd.read_csv(
        PROCESSED_DIR / "charged_off_loss_proxy.csv",
        usecols=["ead_proxy", "lgd_proxy"],
    )
    loss_metrics = pd.read_csv(
        PROCESSED_DIR / "test_with_loss_metrics.csv",
        usecols=[
            "loan_status",
            "actual_default_12m",
            "lgd_proxy",
            "expected_lgd",
            "actual_net_loss",
            "stage2_model_name",
            "stage2_champion_pd_12m",
            "stage2_champion_pd_lifetime",
            "pd_12m_fixed_horizon",
            "lifetime_pd_hazard",
        ],
    )
    return {
        "runtime": runtime,
        "stage2": stage2,
        "view_summary": view_summary,
        "search": search,
        "portfolio": portfolio,
        "segment": segment,
        "hazard_comparison": hazard_comparison,
        "predictions": predictions,
        "charged_proxy": charged_proxy,
        "loss_metrics": loss_metrics,
    }


def _compile_metrics(inputs):
    predictions = inputs["predictions"]
    test_1m = _metrics(predictions["actual_default"], predictions["predicted_hazard_1m"])
    test_12m = _metrics(predictions["actual_default_12m"], predictions["predicted_pd_12m"])
    test_lifetime = _metrics(predictions["actual_default_12m"], predictions["predicted_pd"])

    search = inputs["search"]
    xgb_best = search[search["model_key"].eq("xgboost")].sort_values(
        ["validation_lifetime_auc", "validation_auc"],
        ascending=[False, False],
        na_position="last",
    ).iloc[0]
    rf_best = search[search["model_key"].eq("random_forest")].sort_values(
        ["validation_lifetime_auc", "validation_auc"],
        ascending=[False, False],
        na_position="last",
    ).iloc[0]

    charged_proxy = inputs["charged_proxy"]
    loss_metrics = inputs["loss_metrics"]
    charged_holdout = loss_metrics[
        loss_metrics["loan_status"].eq("Charged Off")
        & loss_metrics["lgd_proxy"].notna()
        & loss_metrics["expected_lgd"].notna()
    ].copy()
    lgd_error = charged_holdout["expected_lgd"] - charged_holdout["lgd_proxy"]
    stage2_model_counts = loss_metrics["stage2_model_name"].fillna("missing").value_counts()

    return {
        "test_1m": test_1m,
        "test_12m": test_12m,
        "test_lifetime": test_lifetime,
        "xgb_best": xgb_best,
        "rf_best": rf_best,
        "charged_proxy_rows": int(len(charged_proxy)),
        "charged_proxy_lgd_mean": float(charged_proxy["lgd_proxy"].mean()),
        "charged_proxy_ead_mean": float(charged_proxy["ead_proxy"].mean()),
        "holdout_lgd_mae": float(lgd_error.abs().mean()),
        "holdout_lgd_rmse": float(math.sqrt((lgd_error**2).mean())),
        "actual_loss_total": float(loss_metrics["actual_net_loss"].fillna(0).sum()),
        "actual_12m_loss_total": float(
            loss_metrics.loc[
                loss_metrics["actual_default_12m"].fillna(0).astype(int) == 1,
                "actual_net_loss",
            ].fillna(0).sum()
        ),
        "stage2_xgb_rows": int(stage2_model_counts.get("XGBoost", 0)),
        "stage2_missing_rows": int(stage2_model_counts.get("missing", 0)),
        "stage2_mean_12m_pd": float(loss_metrics["pd_12m_fixed_horizon"].mean()),
        "stage2_mean_lifetime_pd": float(loss_metrics["lifetime_pd_hazard"].mean()),
    }


def _clear_document(doc):
    body = doc.element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def _set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _set_cell_text(cell, text, bold=False, font_size=8, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(font_size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _set_cell_margin(cell, margin_twips=45):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side in ["top", "start", "bottom", "end"]:
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin_twips))
        node.set(qn("w:type"), "dxa")


def _set_column_width(cell, width):
    cell.width = Inches(width)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def _set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _set_fixed_table_layout(table):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def _add_table(doc, headers, rows, font_size=8, widths=None, numeric_columns=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    _set_fixed_table_layout(table)
    numeric_columns = set(numeric_columns or [])
    if widths is None:
        widths = [6.7 / len(headers)] * len(headers)
    for idx, header in enumerate(headers):
        _set_column_width(table.rows[0].cells[idx], widths[idx])
        _set_cell_text(table.rows[0].cells[idx], header, bold=True, font_size=font_size)
        _set_cell_shading(table.rows[0].cells[idx], "D9EAF7")
        _set_cell_margin(table.rows[0].cells[idx])
    _set_repeat_table_header(table.rows[0])
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            _set_column_width(cells[idx], widths[idx])
            align = WD_ALIGN_PARAGRAPH.RIGHT if idx in numeric_columns else WD_ALIGN_PARAGRAPH.LEFT
            _set_cell_text(cells[idx], value, font_size=font_size, align=align)
            _set_cell_margin(cells[idx])
    doc.add_paragraph()
    return table


def _add_landscape_section(doc):
    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)
    return section


def _add_portrait_section(doc):
    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    section.orientation = WD_ORIENT.PORTRAIT
    if section.page_width > section.page_height:
        section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    return section


def _add_caption(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.italic = True
    run.font.size = Pt(9)


def _add_figure(doc, path, caption, width=6.0):
    path = Path(path)
    if not path.exists():
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(width))
    _add_caption(doc, caption)


def _add_paragraph(doc, text):
    paragraph = doc.add_paragraph(text)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.08
    return paragraph


def _add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def _top_segment_rows(segment):
    axes = ["grade", "fico_bucket", "purpose_group", "annual_income_band", "term_months"]
    rows = []
    for axis in axes:
        axis_df = segment[
            segment["segment_type"].eq(axis)
            & segment["pd_measure"].eq("lifetime")
        ].copy()
        resolved = axis_df[axis_df["analysis_scope"].eq("resolved_test")]
        active = axis_df[axis_df["analysis_scope"].eq("active_snapshot")]
        res_top = resolved.sort_values("total_el", ascending=False).head(1)
        act_top = active.sort_values("total_el", ascending=False).head(1)
        rows.append(
            [
                axis,
                res_top["segment_value"].iat[0] if not res_top.empty else "NA",
                _fmt_pct(res_top["portfolio_el_share"].iat[0]) if not res_top.empty else "NA",
                act_top["segment_value"].iat[0] if not act_top.empty else "NA",
                _fmt_pct(act_top["portfolio_el_share"].iat[0]) if not act_top.empty else "NA",
            ]
        )
    return rows


def _add_title_page(doc):
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Credit Risk Modeling, Expected Loss, and Portfolio Segmentation")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(
        "Calendar-Time Survival PD, LGD/EAD Proxies, Reserve Analytics, and Management Segmentation"
    ).italic = True
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.add_run(
        "Updated with the calibrated full-stage2 XGBoost calendar-time hazard champion and refreshed loss-reserve results."
    )


def _build_document(doc, inputs, metrics):
    runtime = inputs["runtime"]
    stage2 = inputs["stage2"]
    view = inputs["view_summary"]
    search = inputs["search"]
    portfolio = inputs["portfolio"]
    segment = inputs["segment"]
    hazard = inputs["hazard_comparison"]
    xgb_params = stage2["model_params"][stage2["selected_model_key"]]
    xgb_best = metrics["xgb_best"]
    rf_best = metrics["rf_best"]

    _clear_document(doc)
    for section in doc.sections:
        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)

    _add_title_page(doc)

    doc.add_heading("1. Introduction", level=1)
    _add_paragraph(
        doc,
        "This paper presents a credit-risk workflow that connects probability of default, loss given default, exposure at default, expected loss, and portfolio segmentation. The revised implementation uses a calendar-time survival design for the main PD engine: each account-month snapshot predicts next-month charge-off/default, and monthly hazards are aggregated into 12-month and lifetime PD measures.",
    )
    _add_paragraph(
        doc,
        "The current main result is a time-boxed, production-style training run that reuses the existing calendar survival modeling sample and focuses model search on XGBoost and Random Forest. This keeps the workflow executable within a short decision window while preserving chronological validation and downstream reserve compatibility.",
    )

    doc.add_heading("2. Data and Methodology", level=1)
    _add_paragraph(
        doc,
        "The raw input is the LendingClub accepted-loan file. The PD layer is organized as account-month observations keyed by sample_id and snapshot_month. The loss layer keeps loan-level performance fields for LGD, EAD, and expected-loss construction, but those performance fields are excluded from the PD feature set to avoid leakage.",
    )
    _add_table(
        doc,
        ["Split", "Rows", "Event rows", "Cap / policy"],
        [
            ["Train", f"{view['split_counts']['train']:,}", f"{view['event_counts']['train']:,}", "All positives, negatives up to 3:1"],
            ["Validation", f"{view['split_counts']['validation']:,}", f"{view['event_counts']['validation']:,}", "25,000 stratified rows"],
            ["Test", f"{view['split_counts']['test']:,}", f"{view['event_counts']['test']:,}", "100,000 stratified rows"],
            ["Scoring", f"{view['split_counts']['scoring']:,}", f"{view['event_counts']['scoring']:,}", "25,000 active snapshot rows"],
        ],
        widths=[1.2, 1.2, 1.2, 3.1],
        numeric_columns={1, 2},
    )

    doc.add_heading("3. Probability of Default Modeling", level=1)
    doc.add_heading("3.1 Calendar-Time Survival Design", level=2)
    _add_paragraph(
        doc,
        "The target variable is target_1m, the next-month charge-off/default indicator. A resolved default contributes a single event row at the event snapshot, while earlier exposure months and censored active or fully paid observations contribute non-event exposure. Rows are split by snapshot_month, so the same loan can appear in multiple periods as it seasons over time.",
    )
    _add_paragraph(
        doc,
        "The model uses origination-safe borrower and loan features, calendar-time fields, seasoning variables, scheduled-balance proxies, and controlled interactions. Logistic regression uses a compact calendar profile; Random Forest, XGBoost, and the MLP smoke model use the full calendar hazard profile.",
    )

    doc.add_heading("3.2 Time-Boxed Model Search", level=2)
    _add_paragraph(
        doc,
        f"The run mode is {stage2['training_mode']}. It ran {runtime['candidate_count_ran']} candidates in {runtime['total_seconds']:.1f} seconds before downstream loss processing. Grid selection used fast 1-month hazard scoring, while 12-month and lifetime PD aggregation was limited to the best RF and XGBoost candidates and the final champion.",
    )
    candidate_rows = []
    for _, row in search.iterrows():
        candidate_rows.append(
            [
                row["model_name"],
                "Yes" if bool(row["champion_eligible"]) else "Benchmark",
                _short_params(row["params"]),
                _fmt_num(row["validation_auc"], 4),
                _fmt_num(row["validation_ks"], 4),
                _fmt_num(row["fit_seconds"], 1),
            ]
        )
    _add_table(
        doc,
        ["Model", "Champion eligible", "Key parameters", "1m AUC", "KS", "Fit sec"],
        candidate_rows,
        font_size=7,
        widths=[1.15, 1.1, 2.4, 0.75, 0.65, 0.7],
        numeric_columns={3, 4, 5},
    )

    doc.add_heading("3.3 Champion Selection and Test Results", level=2)
    _add_paragraph(
        doc,
        f"The selected main PD champion is XGBoost with n_estimators={xgb_params['n_estimators']}, max_depth={xgb_params['max_depth']}, and learning_rate={xgb_params['learning_rate']}. XGBoost outperformed the RF candidate on validation lifetime proxy AUC and on 1-month validation AUC. For expected-loss use, monthly XGBoost hazards are calibrated with a validation-fitted logit intercept and then aggregated into 12-month and lifetime PD.",
    )
    _add_table(
        doc,
        ["Metric", "Validation", "Test"],
        [
            ["1m AUC", _fmt_num(xgb_best["validation_auc"], 4), _fmt_num(metrics["test_1m"]["auc"], 4)],
            ["1m KS", _fmt_num(xgb_best["validation_ks"], 4), _fmt_num(metrics["test_1m"]["ks"], 4)],
            ["1m Brier", _fmt_num(xgb_best["validation_brier"], 4), _fmt_num(metrics["test_1m"]["brier"], 4)],
            ["12m AUC", _fmt_num(xgb_best["validation_12m_auc"], 4), _fmt_num(metrics["test_12m"]["auc"], 4)],
            ["Lifetime proxy AUC", _fmt_num(xgb_best["validation_lifetime_auc"], 4), _fmt_num(metrics["test_lifetime"]["auc"], 4)],
        ],
        widths=[2.5, 1.1, 1.1],
        numeric_columns={1, 2},
    )
    _add_figure(doc, FIG_TEMPORAL / "validation_roc_comparison.png", "Figure 1. Validation ROC comparison for the calendar-time PD models.")
    _add_figure(doc, FIG_TEMPORAL / "test_roc_comparison.png", "Figure 2. Holdout test ROC for the final XGBoost champion.")
    _add_figure(doc, FIG_TEMPORAL / "champion_test_ks.png", "Figure 3. Champion test KS curve.")
    _add_figure(doc, FIG_TEMPORAL / "xgboost_feature_importance.png", "Figure 4. XGBoost feature importance.")

    doc.add_heading("4. LGD & EAD Estimation", level=1)
    doc.add_heading("4.1 Proxy Construction", level=2)
    _add_paragraph(
        doc,
        "LGD and EAD are constructed from observed charged-off loan recoveries and amortization-aware exposure proxies. EAD_proxy is the unpaid principal exposure at charge-off. LGD_proxy is one minus net recoveries over EAD, clipped to the unit interval. Active loans use current outstanding principal when available and otherwise fall back to a scheduled-balance proxy.",
    )
    _add_table(
        doc,
        ["Diagnostic", "Value"],
        [
            ["Charged-off proxy rows", f"{metrics['charged_proxy_rows']:,}"],
            ["Mean charged-off LGD proxy", _fmt_pct(metrics["charged_proxy_lgd_mean"], 1)],
            ["Average charged-off EAD proxy", _fmt_money(metrics["charged_proxy_ead_mean"])],
            ["Holdout charged-off LGD MAE", _fmt_num(metrics["holdout_lgd_mae"], 4)],
            ["Holdout charged-off LGD RMSE", _fmt_num(metrics["holdout_lgd_rmse"], 4)],
        ],
        widths=[3.4, 1.9],
        numeric_columns={1},
    )
    _add_figure(doc, FIG_LOSS / "lgd_actual_vs_expected.png", "Figure 5. Expected versus realized LGD for charged-off holdout loans.")
    _add_figure(doc, FIG_LOSS / "ead_distribution.png", "Figure 6. EAD distribution used in loss and reserve analytics.")

    doc.add_heading("5. Expected Loss (Loan & Portfolio Level)", level=1)
    _add_paragraph(
        doc,
        "Expected loss is computed as PD x LGD x EAD. The refreshed reserve workflow uses calibrated XGBoost full-stage2 scoring for all resolved holdout and active snapshot loans. The downstream HistGradientBoosting (HGB) hazard model is retained only as an auxiliary diagnostic benchmark and no longer supplies fallback PD for the main EL tables.",
    )
    calibration = stage2.get("hazard_calibration", {}) or {}
    _add_table(
        doc,
        ["Full-stage2 diagnostic", "Value"],
        [
            ["Stage2 model", "XGBoost"],
            ["Rows using XGBoost PD in resolved holdout", f"{metrics['stage2_xgb_rows']:,}"],
            ["Rows with missing stage2 model", f"{metrics['stage2_missing_rows']:,}"],
            ["Mean calibrated 12m PD", _fmt_pct(metrics["stage2_mean_12m_pd"], 1)],
            ["Mean calibrated lifetime PD", _fmt_pct(metrics["stage2_mean_lifetime_pd"], 1)],
            ["Calibration method", calibration.get("method", "none")],
            ["Calibration intercept shift", _fmt_num(calibration.get("intercept_shift"), 4)],
            ["Validation full event rate", _fmt_pct(calibration.get("validation_full_event_rate"), 2)],
        ],
        font_size=7,
        widths=[3.6, 2.2],
        numeric_columns={1},
    )
    portfolio_rows = []
    for _, row in portfolio.iterrows():
        portfolio_rows.append(
            [
                row["analysis_scope"],
                row["pd_measure"],
                f"{int(row['loan_count']):,}",
                _fmt_pct(row["avg_pd"], 1),
                _fmt_pct(row["avg_lgd"], 1),
                _fmt_money(row["total_el"]),
                _fmt_money(row["actual_loss_amount"]) if pd.notna(row["actual_loss_amount"]) else "NA",
            ]
        )
    _add_table(
        doc,
        ["Scope", "PD horizon", "Loans", "Avg PD", "Avg LGD", "Total EL", "Actual net loss"],
        portfolio_rows,
        font_size=7,
        widths=[1.25, 0.85, 0.85, 0.75, 0.75, 1.1, 1.1],
        numeric_columns={2, 3, 4, 5, 6},
    )
    _add_figure(doc, FIG_LOSS / "portfolio_el_comparison.png", "Figure 7. Portfolio expected-loss comparison by scope and horizon.")
    _add_figure(doc, FIG_LOSS / "el_component_summary.png", "Figure 8. PD, LGD, EAD, and EL component summary.")
    _add_figure(doc, FIG_LOSS / "el_concentration_curve.png", "Figure 9. Expected-loss concentration curve.")

    doc.add_heading("6. Risk Segmentation Analysis", level=1)
    _add_paragraph(
        doc,
        "The segmentation layer converts loan-level risk into management views by grade, FICO bucket, purpose group, income band, term, and vintage. These cuts identify where expected loss concentrates and where monitoring or pricing changes would have the highest impact.",
    )
    _add_table(
        doc,
        ["Axis", "Resolved dominant segment", "Resolved EL share", "Active dominant segment", "Active EL share"],
        _top_segment_rows(segment),
        font_size=7,
        widths=[1.2, 1.8, 1.1, 1.8, 1.1],
        numeric_columns={2, 4},
    )
    _add_figure(doc, FIG_LOSS / "segment_el_by_grade.png", "Figure 10. Lifetime expected loss by grade.")
    _add_figure(doc, FIG_LOSS / "grade_fico_heatmap.png", "Figure 11. Grade-FICO heatmap for 12-month PD.")
    _add_figure(doc, FIG_LOSS / "active_pd_compare.png", "Figure 12. Active snapshot PD comparison.")

    doc.add_heading("7. Risk-Based Pricing Design", level=1)
    _add_paragraph(
        doc,
        "The pricing module can use the updated PD, expected LGD, and EAD estimates to set a minimum risk premium. A practical pricing rule would combine expected loss, funding cost, operating expense, capital cost, and a target margin. The current results provide the empirical PD/LGD/EAD inputs but do not yet optimize pricing elasticity or borrower acceptance.",
    )
    _add_bullets(
        doc,
        [
            "Use predicted_pd_12m for short-horizon pricing and early monitoring.",
            "Use predicted_pd for lifetime reserve and capital-style views.",
            "Apply segmentation outputs to review grade, term, and FICO pricing consistency.",
        ],
    )

    doc.add_heading("8. Capital (Economic Capital) Estimation", level=1)
    _add_paragraph(
        doc,
        "Economic capital can be layered on top of expected loss by modeling unexpected loss around the PD/LGD/EAD distribution. The current pipeline supplies portfolio EL and concentration diagnostics; a next step is to estimate loss quantiles under correlated default scenarios and compare those quantiles with expected loss.",
    )

    doc.add_heading("9. Backtesting and Visualization", level=1)
    _add_paragraph(
        doc,
        "Backtesting compares predicted hazards and aggregated PD with realized outcomes by time period, vintage, and segment. The refreshed plots include ROC curves, reliability diagrams, KS curves, score distributions, feature importance, expected-loss components, and segmentation heatmaps.",
    )
    _add_figure(doc, FIG_TEMPORAL / "test_calibration_comparison.png", "Figure 13. Holdout reliability diagram.")
    _add_figure(doc, FIG_LOSS / "hazard_vintage_calibration.png", "Figure 14. Auxiliary reserve hazard vintage calibration.")

    _add_landscape_section(doc)
    doc.add_heading("Appendix A. Complete Tables", level=1)
    _add_paragraph(
        doc,
        "The appendix keeps the wider result tables in fixed-width landscape layout so the main narrative remains readable while preserving the complete tabular outputs used for the paper.",
    )
    appendix_candidate_rows = []
    for _, row in search.iterrows():
        appendix_candidate_rows.append(
            [
                int(row["candidate_id"]),
                row["model_key"],
                "Yes" if bool(row["champion_eligible"]) else "No",
                row["feature_profile"],
                int(row["feature_count"]),
                _short_params(row["params"]),
                _fmt_num(row["validation_auc"], 4),
                _fmt_num(row["validation_12m_auc"], 4),
                _fmt_num(row["validation_lifetime_auc"], 4),
                _fmt_num(row["fit_seconds"], 1),
            ]
        )
    _add_table(
        doc,
        ["ID", "Model", "Champion", "Profile", "Features", "Params", "1m AUC", "12m AUC", "Lifetime AUC", "Fit sec"],
        appendix_candidate_rows,
        font_size=6,
        widths=[0.35, 0.95, 0.6, 1.1, 0.55, 2.1, 0.6, 0.65, 0.8, 0.55],
        numeric_columns={0, 4, 6, 7, 8, 9},
    )
    hazard_rows = []
    for _, row in hazard.iterrows():
        hazard_rows.append(
            [
                row["model_key"],
                row["model_name"],
                "Yes" if bool(row["selected"]) else "No",
                _fmt_num(row["validation_lifetime_auc"], 4),
                _fmt_num(row["validation_12m_auc"], 4),
                _fmt_num(row["test_lifetime_auc"], 4),
                _fmt_num(row["test_12m_auc"], 4),
                int(row["validation_rank"]),
            ]
        )
    _add_table(
        doc,
        ["Aux model", "Name", "Selected", "Val life AUC", "Val 12m AUC", "Test life AUC", "Test 12m AUC", "Rank"],
        hazard_rows,
        font_size=6,
        widths=[1.2, 1.6, 0.65, 0.75, 0.75, 0.8, 0.8, 0.45],
        numeric_columns={3, 4, 5, 6, 7},
    )
    segment_rows = []
    for _, row in segment.iterrows():
        segment_rows.append(
            [
                row["analysis_scope"],
                row["pd_measure"],
                row["segment_type"],
                row["segment_value"],
                f"{int(row['loan_count']):,}",
                _fmt_pct(row["avg_pd"], 1),
                _fmt_pct(row["avg_lgd"], 1),
                _fmt_money(row["total_el"]),
                _fmt_pct(row["portfolio_el_share"], 1),
            ]
        )
    _add_table(
        doc,
        ["Scope", "Horizon", "Axis", "Segment", "Loans", "Avg PD", "Avg LGD", "Total EL", "EL Share"],
        segment_rows,
        font_size=5.5,
        widths=[1.05, 0.65, 1.2, 1.45, 0.75, 0.55, 0.55, 0.9, 0.65],
        numeric_columns={4, 5, 6, 7, 8},
    )

    _add_portrait_section(doc)
    doc.add_heading("10. Conclusion", level=1)
    _add_paragraph(
        doc,
        "The updated workflow replaces the prior static PD framing with a calendar-time survival engine and selects XGBoost as the current main PD champion under a 30-minute training constraint. The champion achieved a validation 1-month AUC of 0.7138 and a test 1-month AUC of 0.7011. The refreshed loss-reserve layer now uses calibrated full-stage2 XGBoost PD for the complete resolved holdout and active reserve populations, then combines PD with credibility-weighted LGD and EAD to produce 12-month and lifetime expected-loss views.",
    )
    _add_paragraph(
        doc,
        "The main limitation is that the current champion is a time-boxed run rather than an exhaustive full-grid search. The model is therefore suitable as the current reproducible result, while a longer background run can be used later to test whether a broader candidate set materially improves validation lifetime PD performance.",
    )


def _backup(path):
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = path.with_name(f"{path.stem}.bak_{timestamp}{path.suffix}")
    shutil.copy2(path, backup_path)
    return backup_path


def main():
    inputs = _read_inputs()
    metrics = _compile_metrics(inputs)
    backups = []
    for target in DOC_TARGETS:
        target.parent.mkdir(parents=True, exist_ok=True)
        if target.exists():
            backups.append(str(_backup(target)))
            doc = Document(target)
        else:
            doc = Document()
        _build_document(doc, inputs, metrics)
        doc.save(target)

    summary = {
        "targets": [str(path) for path in DOC_TARGETS],
        "backups": backups,
        "selected_model": inputs["stage2"]["selected_model_key"],
        "runtime_seconds": inputs["runtime"]["total_seconds"],
        "candidate_count": inputs["runtime"]["candidate_count_ran"],
        "test_1m_auc": metrics["test_1m"]["auc"],
        "test_12m_auc": metrics["test_12m"]["auc"],
        "test_lifetime_auc": metrics["test_lifetime"]["auc"],
    }
    output_path = PROJECT_ROOT / "output" / "doc" / "credit_risk_paper_generation_summary.json"
    output_path.write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
