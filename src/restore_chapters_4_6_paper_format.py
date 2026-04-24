from __future__ import annotations

import json
import math
import shutil
from datetime import datetime
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


PROJECT_ROOT = Path(__file__).resolve().parents[1]
REPO_ROOT = PROJECT_ROOT.parent
PROCESSED_DIR = PROJECT_ROOT / "data" / "processed"
FIGURE_DIR = PROJECT_ROOT / "reports" / "figures" / "loss_reserve"

TARGET_DOCX = [
    PROJECT_ROOT / "output" / "doc" / "credit_risk_selected_chapters_academic.docx",
    REPO_ROOT / "credit risk paper.docx",
]
REFERENCE_CANDIDATES = [
    REPO_ROOT / "credit risk paper.bak_20260424_132814.docx",
    REPO_ROOT / "credit risk paper.docx",
    PROJECT_ROOT / "output" / "doc" / "credit_risk_selected_chapters_academic.docx",
]
REFERENCE_DOCX = next((path for path in REFERENCE_CANDIDATES if path.exists()), REFERENCE_CANDIDATES[-1])


def _fmt_money(value, decimals=1):
    if value is None or pd.isna(value):
        return ""
    value = float(value)
    if abs(value) >= 1_000_000_000:
        amount = f"{value / 1_000_000_000:.3f}".rstrip("0").rstrip(".")
        return f"${amount}B"
    if abs(value) >= 1_000_000:
        return f"${value / 1_000_000:.{decimals}f}M"
    return f"${value:,.0f}"


def _fmt_ead(value):
    if value is None or pd.isna(value):
        return ""
    return f"${float(value):,.0f}"


def _fmt_pct(value, decimals=2):
    if value is None or pd.isna(value):
        return ""
    return f"{float(value) * 100:.{decimals}f}%"


def _set_paragraph_text(doc, index, text):
    paragraph = doc.paragraphs[index]
    style = paragraph.style
    paragraph.text = text
    paragraph.style = style


def _set_cell_text(cell, text):
    cell.text = str(text)


def _set_cell_margins(cell, margin_twips=35):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side in ("top", "start", "bottom", "end"):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin_twips))
        node.set(qn("w:type"), "dxa")


def _set_row_cannot_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def _format_table(table, widths=None, font_size=8, header_font_size=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if hasattr(table, "allow_autofit"):
        table.allow_autofit = False
    header_font_size = header_font_size or font_size
    for row_index, row in enumerate(table.rows):
        _set_row_cannot_split(row)
        for col_index, cell in enumerate(row.cells):
            if widths and col_index < len(widths):
                cell.width = Inches(widths[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.LEFT if row_index == 0 or col_index == 0 else WD_ALIGN_PARAGRAPH.RIGHT
                )
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.size = Pt(header_font_size if row_index == 0 else font_size)
                    if row_index == 0:
                        run.font.bold = True


def _format_tables(doc):
    compact_specs = {
        0: ([1.05, 1.55, 3.65], 7.2),
        1: ([1.05, 1.75, 1.0], 8.0),
        2: ([0.82, 1.62, 1.72, 2.15], 7.0),
        3: ([1.0, 0.55, 0.67, 0.67, 0.72, 0.67, 0.67, 0.72], 6.6),
        4: ([0.88, 0.78, 0.72, 0.70, 0.70, 0.86, 0.82, 0.92], 7.0),
        5: ([0.45, 1.0, 0.82, 1.0, 0.82], 7.3),
        6: ([1.05, 1.3, 0.75, 1.3, 0.75], 7.0),
    }
    for table_index, (widths, size) in compact_specs.items():
        if table_index < len(doc.tables):
            _format_table(doc.tables[table_index], widths=widths, font_size=size)
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith("Table 6.1."):
            paragraph.paragraph_format.page_break_before = True


def _replace_inline_image(doc, shape_index, image_path):
    blips = doc.element.body.xpath(".//a:blip")
    if shape_index >= len(blips):
        raise IndexError(f"Shape index {shape_index} not found in template.")
    rel_id = blips[shape_index].get(qn("r:embed"))
    image_part = doc.part.related_parts[rel_id]
    image_part._blob = Path(image_path).read_bytes()


def _load_inputs():
    charged = pd.read_csv(PROCESSED_DIR / "charged_off_loss_proxy.csv")
    resolved = pd.read_csv(PROCESSED_DIR / "test_with_loss_metrics.csv")
    active = pd.read_csv(PROCESSED_DIR / "cecl_active_snapshot.csv")
    portfolio = pd.read_csv(PROCESSED_DIR / "portfolio_expected_loss_summary.csv")
    segment = pd.read_csv(PROCESSED_DIR / "segment_expected_loss_summary.csv")
    hazard = pd.read_csv(PROCESSED_DIR / "hazard_model_comparison.csv")
    stage2 = json.loads((PROCESSED_DIR / "stage2_champion_config.json").read_text())
    return {
        "charged": charged,
        "resolved": resolved,
        "active": active,
        "portfolio": portfolio,
        "segment": segment,
        "hazard": hazard,
        "stage2": stage2,
    }


def _metrics(inputs):
    charged = inputs["charged"]
    resolved = inputs["resolved"]
    active = inputs["active"]
    portfolio = inputs["portfolio"]
    segment = inputs["segment"]
    charged_holdout = resolved[
        resolved["loan_status"].eq("Charged Off")
        & resolved["lgd_proxy"].notna()
        & resolved["expected_lgd"].notna()
    ].copy()
    lgd_error = charged_holdout["expected_lgd"] - charged_holdout["lgd_proxy"]

    def portfolio_row(scope, measure):
        return portfolio[
            portfolio["analysis_scope"].eq(scope) & portfolio["pd_measure"].eq(measure)
        ].iloc[0]

    top_active_share = (
        active.nlargest(1000, "lifetime_expected_loss")["lifetime_expected_loss"].sum()
        / active["lifetime_expected_loss"].sum()
    )
    top_resolved_share = (
        resolved.nlargest(1000, "expected_loss_lifetime")["expected_loss_lifetime"].sum()
        / resolved["expected_loss_lifetime"].sum()
    )

    return {
        "ead_mean": charged["ead_proxy"].mean(),
        "ead_median": charged["ead_proxy"].median(),
        "ead_p90": charged["ead_proxy"].quantile(0.9),
        "charged_lgd_mean": charged["lgd_proxy"].mean(),
        "holdout_lgd_mean": charged_holdout["lgd_proxy"].mean(),
        "holdout_lgd_mae": lgd_error.abs().mean(),
        "holdout_lgd_rmse": math.sqrt((lgd_error**2).mean()),
        "portfolio_rows": {
            ("observable_12m_test", "12m"): portfolio_row("observable_12m_test", "12m"),
            ("resolved_test", "12m"): portfolio_row("resolved_test", "12m"),
            ("resolved_test", "lifetime"): portfolio_row("resolved_test", "lifetime"),
            ("active_snapshot", "12m"): portfolio_row("active_snapshot", "12m"),
            ("active_snapshot", "lifetime"): portfolio_row("active_snapshot", "lifetime"),
        },
        "top_active_share": top_active_share,
        "top_resolved_share": top_resolved_share,
        "segment": segment,
    }


def _update_tables(doc, inputs, metrics):
    resolved = inputs["resolved"]
    active = inputs["active"]
    hazard = inputs["hazard"]
    segment = inputs["segment"]

    lookup = (
        pd.concat(
            [
                resolved[["sample_id", "lgd_lookup_source"]].assign(scope="Resolved holdout"),
                active[["sample_id", "lgd_lookup_source"]].assign(scope="Active snapshot"),
            ]
        )
        .groupby(["scope", "lgd_lookup_source"], as_index=False)
        .agg(loan_count=("sample_id", "count"))
    )
    lookup_map = {
        (row.scope, row.lgd_lookup_source): int(row.loan_count)
        for row in lookup.itertuples(index=False)
    }
    table = doc.tables[1]
    rows = [
        ("Active snapshot", "Exact segment", lookup_map.get(("Active snapshot", "exact_segment"), 0)),
        ("Active snapshot", "Grade-term fallback", lookup_map.get(("Active snapshot", "grade_term"), 0)),
        ("Resolved holdout", "Grade-term fallback", lookup_map.get(("Resolved holdout", "grade_term"), 0)),
    ]
    for idx, row in enumerate(rows, start=1):
        for col, value in enumerate(row):
            _set_cell_text(table.rows[idx].cells[col], f"{value:,}" if col == 2 else value)

    table = doc.tables[2]
    mapping_updates = {
        1: [
            "12-month PD",
            "Probability that a loan defaults within the next twelve months",
            "Direct active-snapshot XGBoost model",
            "Near-term expected-loss measurement for active-at-snapshot cohorts",
        ],
        2: [
            "Lifetime PD",
            "Probability that a loan defaults over its remaining contractual life",
            "Original static HistGradientBoosting loan-level model",
            "Remaining-life reserve-style analytics",
        ],
    }
    for row_index, values in mapping_updates.items():
        for col, value in enumerate(values):
            _set_cell_text(table.rows[row_index].cells[col], value)

    table = doc.tables[3]
    for row_index, (_, row) in enumerate(hazard.iterrows(), start=1):
        values = [
            row["model_name"],
            "True" if bool(row["selected"]) else "False",
            f"{row['validation_lifetime_auc']:.4f}",
            f"{row['validation_lifetime_ks']:.4f}",
            f"{row['validation_lifetime_brier']:.4f}",
            f"{row['test_lifetime_auc']:.4f}",
            f"{row['test_lifetime_ks']:.4f}",
            f"{row['test_lifetime_brier']:.4f}",
        ]
        for col, value in enumerate(values):
            _set_cell_text(table.rows[row_index].cells[col], value)

    table = doc.tables[4]
    order = [
        ("Full 12M cohort", "12-month", metrics["portfolio_rows"][("observable_12m_test", "12m")]),
        ("Resolved holdout", "Lifetime", metrics["portfolio_rows"][("resolved_test", "lifetime")]),
        ("Active snapshot", "12-month", metrics["portfolio_rows"][("active_snapshot", "12m")]),
        ("Active snapshot", "Lifetime", metrics["portfolio_rows"][("active_snapshot", "lifetime")]),
    ]
    for row_index, (scope, horizon, row) in enumerate(order, start=1):
        values = [
            scope,
            horizon,
            f"{int(row['loan_count']):,}",
            f"{row['avg_pd'] * 100:.2f}",
            f"{row['avg_lgd'] * 100:.2f}",
            _fmt_ead(row["avg_ead"]),
            _fmt_money(row["total_el"]),
            _fmt_money(row["actual_loss_amount"]) if pd.notna(row["actual_loss_amount"]) else "",
        ]
        for col, value in enumerate(values):
            _set_cell_text(table.rows[row_index].cells[col], value)

    grade = segment[
        segment["segment_type"].eq("grade") & segment["pd_measure"].eq("lifetime")
    ].copy()
    table = doc.tables[5]
    for row_index, grade_value in enumerate(list("ABCDEFG"), start=1):
        resolved_row = grade[
            grade["analysis_scope"].eq("resolved_test")
            & grade["segment_value"].astype(str).eq(grade_value)
        ].iloc[0]
        active_row = grade[
            grade["analysis_scope"].eq("active_snapshot")
            & grade["segment_value"].astype(str).eq(grade_value)
        ].iloc[0]
        values = [
            grade_value,
            _fmt_money(resolved_row["total_el"]),
            _fmt_pct(resolved_row["portfolio_el_share"]),
            _fmt_money(active_row["total_el"]),
            _fmt_pct(active_row["portfolio_el_share"]),
        ]
        for col, value in enumerate(values):
            _set_cell_text(table.rows[row_index].cells[col], value)

    axis_labels = [
        ("Credit grade", "grade"),
        ("Loan term", "term_months"),
        ("Loan purpose", "purpose_group"),
        ("FICO bucket", "fico_bucket"),
        ("Income band", "annual_income_band"),
    ]
    table = doc.tables[6]
    for row_index, (label, axis) in enumerate(axis_labels, start=1):
        resolved_row = segment[
            segment["segment_type"].eq(axis)
            & segment["pd_measure"].eq("lifetime")
            & segment["analysis_scope"].eq("resolved_test")
        ].sort_values("total_el", ascending=False).iloc[0]
        active_row = segment[
            segment["segment_type"].eq(axis)
            & segment["pd_measure"].eq("lifetime")
            & segment["analysis_scope"].eq("active_snapshot")
        ].sort_values("total_el", ascending=False).iloc[0]
        values = [
            label,
            _segment_label(axis, resolved_row["segment_value"]),
            _fmt_pct(resolved_row["portfolio_el_share"]),
            _segment_label(axis, active_row["segment_value"]),
            _fmt_pct(active_row["portfolio_el_share"]),
        ]
        for col, value in enumerate(values):
            _set_cell_text(table.rows[row_index].cells[col], value)


def _segment_label(axis, value):
    if axis == "term_months":
        return f"{int(float(value))} months"
    if axis == "annual_income_band":
        mapping = {"50-100k": "$50,000-$100,000", "<50k": "<$50,000", "100-150k": "$100,000-$150,000", "150k+": "$150,000+"}
        return mapping.get(str(value), str(value))
    if axis in {"purpose_group", "fico_bucket"}:
        return str(value).replace("_", " ").title()
    return str(value)


def _update_paragraphs(doc, inputs, metrics):
    p = metrics["portfolio_rows"]
    observable_12 = p[("observable_12m_test", "12m")]
    resolved_12 = p[("resolved_test", "12m")]
    resolved_life = p[("resolved_test", "lifetime")]
    active_12 = p[("active_snapshot", "12m")]
    active_life = p[("active_snapshot", "lifetime")]
    gap12 = observable_12["total_el"] - observable_12["actual_loss_amount"]
    gaplife = resolved_life["actual_loss_amount"] - resolved_life["total_el"]
    segment = metrics["segment"]

    grade = segment[
        segment["segment_type"].eq("grade") & segment["pd_measure"].eq("lifetime")
    ].copy()

    def grade_share(scope, grade_value):
        row = grade[
            grade["analysis_scope"].eq(scope)
            & grade["segment_value"].astype(str).eq(grade_value)
        ].iloc[0]
        return row["portfolio_el_share"]

    def leading_axis(scope, axis):
        row = (
            segment[
                segment["analysis_scope"].eq(scope)
                & segment["pd_measure"].eq("lifetime")
                & segment["segment_type"].eq(axis)
            ]
            .sort_values("total_el", ascending=False)
            .iloc[0]
        )
        return _segment_label(axis, row["segment_value"]), row["portfolio_el_share"]

    resolved_purpose, resolved_purpose_share = leading_axis("resolved_test", "purpose_group")
    active_purpose, active_purpose_share = leading_axis("active_snapshot", "purpose_group")
    resolved_term, resolved_term_share = leading_axis("resolved_test", "term_months")
    active_term, active_term_share = leading_axis("active_snapshot", "term_months")

    _set_paragraph_text(
        doc,
        2,
        "Chapters 4 through 6 are updated below to reflect the dual-PD framework, refreshed loss-reserve application, and segmentation results.",
    )
    _set_paragraph_text(
        doc,
        17,
        f"In the charged-off proxy sample, the mean exposure proxy equals ${metrics['ead_mean'] / 1000:.1f} thousand, the median equals ${metrics['ead_median'] / 1000:.1f} thousand, and the 90th percentile equals ${metrics['ead_p90'] / 1000:.1f} thousand. These magnitudes confirm that default severity in this portfolio is economically material even after partial principal recovery.",
    )
    _set_paragraph_text(
        doc,
        32,
        f"The charged-off proxy sample has a mean realized LGD of {_fmt_pct(metrics['charged_lgd_mean'])}, which already suggests that unsecured installment defaults in this dataset are characterized by low recovery. On the temporally held-out charged-off test loans, the credibility-weighted expected LGD lookup produces an MAE of {metrics['holdout_lgd_mae']:.4f} and an RMSE of {metrics['holdout_lgd_rmse']:.4f}.",
    )
    _set_paragraph_text(
        doc,
        41,
        "Expected loss is the bridge between borrower-level risk estimation and portfolio-level financial measurement. The updated framework now uses the original static HistGradientBoosting loan-level model for lifetime PD and a direct active-snapshot XGBoost model for twelve-month PD.",
    )
    _set_paragraph_text(
        doc,
        44,
        "The auxiliary reserve hazard comparison remains useful for governance because it documents the intermediate monthly survival approach. It is no longer the main twelve-month PD input: lifetime expected loss uses the rerun original static model, while twelve-month expected loss uses the calibrated direct active-snapshot model.",
    )
    _set_paragraph_text(
        doc,
        46,
        "Table 5.2. Auxiliary legacy reserve hazard benchmark, not used as the main PD input.",
    )
    _set_paragraph_text(
        doc,
        48,
        "The main twelve-month PD model directly estimates whether an active-at-snapshot loan will charge off or default during the next twelve months. Each row represents an account-month snapshot identified by sample_id and snapshot_month, restricted to loans still at risk at that snapshot. Splits are based on snapshot_month rather than issue_date, so the design approximates a bank's recurring monthly production risk file.",
    )
    _set_paragraph_text(
        doc,
        49,
        "The earlier one-month hazard model is retained as a comparison group and research diagnostic. It converts monthly hazards into cumulative twelve-month PDs by summing log survival terms, but the final twelve-month EL input now comes from the direct active-snapshot XGBoost model because that target is aligned with the intended current-portfolio twelve-month decision horizon.",
    )
    _set_paragraph_text(
        doc,
        50,
        "PD_12m = calibrated direct active-snapshot probability of default within the next twelve months;   PD_life = static HGB calibrated loan-level probability.",
    )
    _set_paragraph_text(
        doc,
        51,
        "Equation (5.1). Direct twelve-month PD and static lifetime PD mapping.",
    )
    _set_paragraph_text(
        doc,
        54,
        "The distinction between horizons is equally important. Twelve-month expected loss uses calibrated direct active-snapshot predicted_pd_12m, whereas lifetime expected loss uses the rerun original static HistGradientBoosting predicted_pd. This two-model design avoids forcing one specification to serve both short-horizon monitoring and lifetime reserve analytics.",
    )
    _set_paragraph_text(
        doc,
        63,
        f"On the full observable twelve-month cohort, the model-implied expected loss equals {_fmt_money(observable_12['total_el'])}, whereas realized twelve-month net loss equals {_fmt_money(observable_12['actual_loss_amount'])}. The conservative overlay therefore makes expected loss exceed realized loss by {_fmt_money(gap12)}, and the expected-loss coverage ratio is {observable_12['total_el'] / observable_12['actual_loss_amount'] * 100:.2f}% of realized twelve-month loss. The resolved-only twelve-month comparison is retained as a diagnostic because that subset excludes surviving active loans and therefore overstates default intensity.",
    )
    _set_paragraph_text(
        doc,
        64,
        f"For the active snapshot, the reported twelve-month EL equals {_fmt_money(active_12['total_el'])}, while the lifetime reserve-style EL equals {_fmt_money(active_life['total_el'])}. The lifetime result is based on the static HGB model and is higher because its average lifetime PD ({_fmt_pct(active_life['avg_pd'])}) exceeds the average direct twelve-month PD ({_fmt_pct(active_12['avg_pd'])}).",
    )
    _set_paragraph_text(
        doc,
        69,
        "Figure 5.2 reports expected and realized loss separately for the full observable twelve-month cohort and the resolved lifetime diagnostic.",
    )
    _set_paragraph_text(
        doc,
        71,
        "Figure 5.2. Expected and realized loss by horizon under the dual-PD framework.",
    )
    _set_paragraph_text(
        doc,
        72,
        f"The evidence indicates that the full-cohort twelve-month comparison is materially less distorted than the resolved-only diagnostic. For lifetime, predicted EL equals {_fmt_money(resolved_life['total_el'])} versus realized lifetime net loss of {_fmt_money(resolved_life['actual_loss_amount'])}, covering about {resolved_life['total_el'] / resolved_life['actual_loss_amount'] * 100:.2f}% of realized loss.",
    )
    _set_paragraph_text(
        doc,
        75,
        f"Concentration diagnostics reinforce this point. The top 1,000 loans explain only {_fmt_pct(metrics['top_active_share'])} of lifetime EL in the active snapshot and {_fmt_pct(metrics['top_resolved_share'])} in the resolved holdout. Expected loss in this portfolio is therefore not dominated by a tiny number of single-name exposures.",
    )
    _set_paragraph_text(
        doc,
        90,
        f"Several concentration patterns are immediate. In the active snapshot, grade C alone contributes {_fmt_pct(grade_share('active_snapshot', 'C'))} of lifetime EL, followed by grade D at {_fmt_pct(grade_share('active_snapshot', 'D'))} and grade B at {_fmt_pct(grade_share('active_snapshot', 'B'))}. The resolved holdout exhibits a similar but slightly less concentrated pattern, with grade C contributing {_fmt_pct(grade_share('resolved_test', 'C'))} and grade D contributing {_fmt_pct(grade_share('resolved_test', 'D'))}.",
    )
    _set_paragraph_text(
        doc,
        91,
        f"The cross-axis concentration statistics reinforce the same story. {resolved_purpose} loans account for {_fmt_pct(resolved_purpose_share)} of resolved lifetime EL, while {active_purpose} loans account for {_fmt_pct(active_purpose_share)} of active lifetime EL. The leading term segment is {resolved_term} in the resolved holdout ({_fmt_pct(resolved_term_share)}) and {active_term} in the active snapshot ({_fmt_pct(active_term_share)}), indicating that contractual tenor remains a primary driver of reserve exposure.",
    )
    _set_paragraph_text(
        doc,
        95,
        "The grade comparison makes clear that the active portfolio is not simply a scaled-up version of the resolved holdout. It remains concentrated in grades B through D, especially grade C, while the weakest grades carry high loan-level severity but smaller aggregate shares because fewer loans sit in those buckets.",
    )


def _replace_figures(doc):
    replacements = {
        1: FIGURE_DIR / "ead_by_grade.png",
        4: FIGURE_DIR / "lgd_actual_vs_expected.png",
        7: FIGURE_DIR / "portfolio_el_comparison.png",
        8: FIGURE_DIR / "predicted_vs_actual_loss.png",
        10: FIGURE_DIR / "segment_el_by_grade.png",
        11: FIGURE_DIR / "grade_term_heatmap.png",
    }
    for shape_index, path in replacements.items():
        _replace_inline_image(doc, shape_index, path)


def _backup(path):
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup = path.with_name(f"{path.stem}.bak_{timestamp}{path.suffix}")
    if path.exists():
        shutil.copy2(path, backup)
        return str(backup)
    return None


def build_document(target, inputs, metrics):
    doc = Document(REFERENCE_DOCX)
    _replace_figures(doc)
    _update_paragraphs(doc, inputs, metrics)
    _update_tables(doc, inputs, metrics)
    _format_tables(doc)
    target.parent.mkdir(parents=True, exist_ok=True)
    backup = _backup(target)
    doc.save(target)
    return backup


def main():
    inputs = _load_inputs()
    metrics = _metrics(inputs)
    backups = []
    for target in TARGET_DOCX:
        backups.append(build_document(target, inputs, metrics))
    summary = {
        "template": str(REFERENCE_DOCX),
        "targets": [str(path) for path in TARGET_DOCX],
        "backups": backups,
        "scope": "chapters_4_to_6_only",
        "observable_12m_el": float(metrics["portfolio_rows"][("observable_12m_test", "12m")]["total_el"]),
        "resolved_12m_diagnostic_el": float(metrics["portfolio_rows"][("resolved_test", "12m")]["total_el"]),
        "resolved_lifetime_el": float(metrics["portfolio_rows"][("resolved_test", "lifetime")]["total_el"]),
        "active_lifetime_el": float(metrics["portfolio_rows"][("active_snapshot", "lifetime")]["total_el"]),
    }
    output_path = PROJECT_ROOT / "output" / "doc" / "chapters_4_6_restore_summary.json"
    output_path.write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
